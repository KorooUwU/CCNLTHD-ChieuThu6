from pathlib import Path
import re, json, subprocess, textwrap
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
WORK = Path(__file__).parent
OUT = ROOT / 'outputs' / 'report-authentication'
OUT.mkdir(parents=True, exist_ok=True)
SRC = Path('D:/Study/CCNLTHD/BÁO CÁO ĐỒ ÁN.docx')
doc = Document(SRC)
assets = WORK / 'assets'
assets.mkdir(exist_ok=True)
def read_log(path):
    data=path.read_bytes()
    return data.decode('utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig', errors='replace')
startup = read_log(ROOT/'notification-dispatcher/auth-startup.log')
verify = read_log(ROOT/'notification-dispatcher/auth-verify.log')
commit = subprocess.check_output(['git','rev-parse','--short','HEAD'], cwd=ROOT, text=True).strip()

def image_log(name, title, lines):
    font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', 23)
    title_font = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 23)
    wrapped=[]
    for line in lines:
        wrapped.extend(textwrap.wrap(line, width=75, replace_whitespace=False) or [''])
    im=Image.new('RGB',(1100,100+len(wrapped)*30),(248,249,251))
    draw=ImageDraw.Draw(im)
    draw.text((24,18),title,fill='black',font=title_font)
    for i,line in enumerate(wrapped): draw.text((24,68+i*30),line,fill=(25,32,42),font=font)
    path=assets/(name+'.png'); im.save(path)
    return path

env_lines=['JAVA 21 / MAVEN WRAPPER - kiem tra ngay 15/10/2026']
env_lines += subprocess.run(['java','-version'],capture_output=True,text=True).stderr.splitlines()[:3]
env_lines += subprocess.run(['cmd','/c',str(ROOT/'notification-dispatcher/mvnw.cmd'),'-v'],capture_output=True,text=True).stdout.splitlines()[:1]
env_img=image_log('environment','MC01 - Trich ket qua lenh kiem tra moi truong',env_lines)
test_lines=[line for line in verify.splitlines() if 'Tests run:' in line or 'BUILD SUCCESS' in line or 'Finished at:' in line]
test_img=image_log('tests','MC04 - Trich log Maven thuc te ngay 14/10/2026',test_lines)
db_lines=[]
for line in startup.splitlines():
    if 'Database JDBC URL' in line or 'Database version:' in line:
        db_lines.append(line.strip())
    if 'Successfully applied 3 migrations' in line:
        db_lines.append(line[line.index('Successfully applied'):])
    if 'Started NotificationDispatcherApplication in' in line:
        db_lines.append(line[line.index('Started NotificationDispatcherApplication'):])
db_img=image_log('database','MC02 - PostgreSQL, Flyway va khoi dong ung dung',db_lines)
sql_start=startup.index('Hibernate:')
sql_lines=startup[sql_start:].splitlines()[:25]
sql_img=image_log('sql','MC03 - SQL Hibernate trong luong dang ky',sql_lines)
cond_lines=verify.splitlines()
selected=[]
for key in ['DataSourceAutoConfiguration matched:', 'DataSourceConfiguration.Hikari matched:', 'HibernateJpaAutoConfiguration matched:']:
    for i,line in enumerate(cond_lines):
        if key in line:
            selected += cond_lines[i:i+4] + ['']
            break
cond_img=image_log('conditions','MC05 - Trich Conditions Evaluation Report',selected)

def para(text='', style=None):
    p=doc.add_paragraph(style=style or 'Normal')
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=False;r.font.cs_bold=False
    return p

def heading(text,level=4):
    p=para(text,f'Heading {level}')
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(12)
    for r in p.runs:r.font.bold=True;r.font.cs_bold=True;r.font.color.rgb=RGBColor(0,0,0);r.font.size=Pt(14 if level<=3 else 13)
    return p

def code(text):
    p=para(text.strip())
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.left_indent=Cm(.15)
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.right_indent=Cm(0)
    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8.5)
    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F3F4F6');p._p.get_or_add_pPr().append(shading)
    return p

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'A0A0A0');borders.append(edge)
    t._tbl.tblPr.append(borders)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        for cell,val in zip(t.add_row().cells,row):cell.text=str(val)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0);p.paragraph_format.first_line_indent=Cm(0)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ri==0;r.font.cs_bold=ri==0
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF2');cell._tc.get_or_add_tcPr().append(sh)
    return t

def picture(path,caption):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(15.3))
    cap=para(caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:r.font.italic=True;r.font.size=Pt(10)

def anchor(text):
    return next(p for p in doc.paragraphs if p.text.strip()==text)

def insert_block_before(anchor_p, fn):
    before=list(doc.element.body)
    fn()
    additions=[e for e in list(doc.element.body) if e not in before and e.tag!=qn('w:sectPr')]
    for e in additions:anchor_p._p.addprevious(e)

def replace_section(start,end,fn):
    a=anchor(start);b=anchor(end)
    e=a._p.getnext()
    while e is not None and e is not b._p:
        nxt=e.getnext();e.getparent().remove(e);e=nxt
    insert_block_before(b,fn)

# Replace handwritten Chat API illustration with an implemented request flow.
def section32():
    para('Nhóm sử dụng chức năng đăng ký tài khoản làm ví dụ đã triển khai. Request JSON đi qua chuỗi Servlet Filter/Spring Security, DispatcherServlet và Bean Validation trước khi AuthController gọi AuthService. Service kiểm tra dữ liệu trùng, mã hóa mật khẩu và gọi TaiKhoanRepository; PostgreSQL là nơi thực thi ràng buộc UNIQUE cuối cùng.')
    code('Client -> SecurityFilterChain -> DispatcherServlet -> @Valid\n  -> AuthController.register -> proxy @Transactional\n  -> AuthService.register -> TaiKhoanRepository.saveAndFlush\n  -> Hibernate SQL -> PostgreSQL -> UserResponseDTO -> HTTP 201')
    para('Đây là kiến trúc phân tầng cho REST, không phải backend sinh trang View HTML. DTO quyết định JSON trao đổi, còn Entity mô tả ánh xạ dữ liệu. Code đối chiếu: controller/AuthController.java, service/AuthService.java và repository/TaiKhoanRepository.java trong src/main/java/com/notification/dispatcher/.')
    heading('3.2.3. Kết quả minh họa',4)
    para('Phiên kiểm thử ngày 14/10/2026 trên PostgreSQL tạm riêng đã xác nhận đăng ký trả 201, đăng nhập trả 200 và /api/users/me trả 200 với đúng chủ token. Đăng ký trùng trả 409; sai mật khẩu hoặc thiếu/sai token trả 401. Bộ sưu tập Postman chạy bằng Newman gồm 9 request và 11 assertion, không lỗi. Các kết quả này chỉ chứng minh Authentication; API lịch sử chat, SSE và RabbitMQ chưa được coi là đã hoàn thành.')
    table(['API','Đầu vào hoặc xác thực','Kết quả thực nghiệm'],[
        ('POST /api/auth/register','tenTaiKhoan, email, userName, password','201; trả id UUID, tên, email, username, createdAt; không có password/passwordHash'),
        ('POST /api/auth/login','userName và password','200; accessToken, tokenType=Bearer, expiresIn=900 mặc định'),
        ('GET /api/users/me','Authorization: Bearer <token>','200; danh tính lấy từ sub của token, không từ userId trong query')
    ],[4,5.1,6.8])
    para('Bản ghi API không chứa token hay mật khẩu thật trong báo cáo. Collection tái hiện nằm tại docs/postman/authentication.postman_collection.json; các giả định và giới hạn được trình bày ở Chương 5.')

replace_section('3.2.2. Minh họa luồng xử lý yêu cầu (MVC và kiến trúc phân tầng)',
                '3.3. Minh chứng kiến trúc dự án', section32)

def section34():
    para(f'Mã nguồn đối chiếu tại commit {commit} (14/10/2026, “hoan thien authentication”). Log thực nghiệm được lưu ngày 14/10/2026; phiên bản công cụ được đối chiếu lại ngày 15/10/2026. Bảng sau ghi kết quả đã quan sát, không thay thế kết quả thực tế bằng mục tiêu mong đợi.')
    table(['Hạng mục','Cách kiểm tra','Kết quả đã quan sát','Minh chứng'],[
        ('JDK, Maven, Spring Boot','java -version; .\\mvnw.cmd -v; pom.xml','Java 21.0.8; Maven Wrapper 3.9.16; Spring Boot 4.1.1','MC01; pom.xml'),
        ('Ứng dụng khởi động','spring-boot:run, cổng test 18081','Started NotificationDispatcherApplication; PostgreSQL tạm tại 15432','MC02; auth-startup.log, 14/10'),
        ('PostgreSQL và schema','Kết nối JPA và Flyway trên DB trống','PostgreSQL 17.11; ba migration được áp dụng, schema v3','MC02; V1–V3'),
        ('Luồng Authentication','Collection Postman qua Newman','9 request, 11 assertion thành công; chi tiết ở 5.5','Collection; biên bản kiểm thử 14/10'),
        ('Build và test','clean verify','29 test; Failures=0, Errors=0; BUILD SUCCESS; tạo JAR','MC04; auth-verify.log'),
        ('Auto-configuration','Conditions Evaluation Report','DataSource/Hikari và Hibernate JPA có điều kiện matched','MC05; 4.4'),
    ],[2.8,4.1,5.2,3.8])
    picture(env_img,'Minh chứng MC01. Trích kết quả kiểm tra công cụ, ngày 15/10/2026.')
    picture(db_img,'Minh chứng MC02. Trích log khởi động và migration thực tế ngày 14/10/2026.')
    para('Cổng 18081 và 15432 dùng để cô lập kiểm thử, không phải thay đổi cổng mặc định 8081 của ứng dụng. Database/container tạm đã được dọn sau phiên chạy. Không có số liệu tải hoặc độ trễ tổng quát được suy ra từ thời gian khởi động này.')

replace_section('3.4. Thực hành khởi tạo / kiểm chứng môi trường','Chương 4: Bắt buộc lõi của Spring Boot',section34)

def chapter4():
    para('Chương này giải thích cơ chế Spring bằng các thành phần đang chạy trong Authentication: AuthController, AuthService, TaiKhoanRepository, PasswordEncoder, JwtService và cấu hình JPA/Flyway. Mỗi mục phân biệt code đang có, kết quả đã quan sát và tình huống chưa kiểm chứng.')
    heading('4.1 IoC và Dependency Injection',4)
    heading('4.1.1 Bean và constructor injection trong dự án',5)
    para('Trong quá trình ApplicationContext refresh, component scan đăng ký AuthController (@RestController), AuthService và JwtService (@Service). SecurityConfig khai báo PasswordEncoder bằng @Bean; Spring Data tạo proxy triển khai TaiKhoanRepository. Spring giải quyết các dependency khi tạo bean. Với cấu hình mặc định, các singleton không lazy thường được tạo trong giai đoạn khởi động; thứ tự giữa các bean phụ thuộc dependency, không chỉ thứ tự file.')
    code('public AuthService(TaiKhoanRepository taiKhoanRepository,\n        PasswordEncoder passwordEncoder, JwtService jwtService) {\n    this.taiKhoanRepository = taiKhoanRepository;\n    this.passwordEncoder = passwordEncoder;\n    this.jwtService = jwtService;\n}')
    table(['Thành phần','Cách được tạo','Scope và vai trò'],[
        ('AuthService','@Service, constructor injection','Singleton mặc định; nghiệp vụ đăng ký/đăng nhập'),
        ('TaiKhoanRepository','Spring Data tạo repository proxy','Singleton bean; thực thi JPA query method'),
        ('PasswordEncoder','@Bean trong SecurityConfig','Singleton bean BCryptPasswordEncoder'),
        ('JwtEncoder/JwtDecoder','@Bean trong JwtConfig','Singleton; ký và kiểm chứng token'),
    ],[3.2,5.3,7.4])
    para('Singleton ở đây là một instance cho mỗi bean definition trong một ApplicationContext, không phải singleton toàn JVM. Nhiều request có thể dùng cùng AuthService; service không lưu thông tin người dùng vào biến trạng thái dùng chung. Ví dụ scope khác gồm prototype, request và session; dự án chưa cần thêm các scope này.')
    heading('4.1.2 Hai cài đặt và cách chọn dependency',5)
    para('Code hiện chỉ có một PasswordEncoder bean nên injection không mơ hồ. Nếu có hai bean cùng kiểu, có thể chọn mặc định bằng @Primary hoặc chỉ định bean tại điểm inject bằng @Qualifier. Ví dụ sau là tình huống giải thích, không phải cấu hình đang bật:')
    code('@Bean @Primary\nPasswordEncoder defaultEncoder() {\n    return new BCryptPasswordEncoder();\n}\n// Neu co bean thu hai: dung @Qualifier("defaultEncoder")\n// tai tham so constructor de chi dinh ro.')
    para('Không thêm bean thứ hai chỉ để đủ annotation. Nếu không có ứng viên ưu tiên hoặc chỉ định rõ khi nhiều bean phù hợp, việc khởi tạo phụ thuộc có thể thất bại vì không chọn được bean duy nhất.')
    heading('4.1.3 Vòng đời bean và hậu quả của tự new',5)
    para('Vòng đời khái quát: tạo instance → cung cấp dependency → callback khởi tạo (như @PostConstruct) → bean post-processing và proxy khi cần → sử dụng → callback hủy khi context đóng. Trong dự án, @Bean(initMethod="migrate") ở FlywayConfig yêu cầu Spring gọi migrate() khi khởi tạo Flyway; MC02 xác nhận ba migration đã chạy. Không mô tả đây là log @PostConstruct vì code chưa có callback đó.')
    para('Ví dụ tự kiểm có thể thêm tạm @PostConstruct để ghi log khởi tạo và @PreDestroy để ghi log đóng context; kết quả callback chưa được thu trong bộ minh chứng hiện tại. Callback hủy của singleton được container quản lý khi đóng bình thường; không bảo đảm khi tiến trình bị kill cưỡng bức. Spring không tự quản lý toàn bộ việc hủy các prototype bean.')
    para('Nếu Controller tự new AuthService(...), đối tượng đó không đi qua container/proxy như bean AuthService; @Transactional và dependency/callback do Spring quản lý sẽ không tự áp dụng chỉ vì lớp có annotation. Ngược lại, new AuthService(mockRepository, encoder, mockJwt) trong unit test là có chủ đích để kiểm tra logic thuần; test đó không chứng minh transaction của Spring.')
    heading('4.2 REST Controller và tầng Service',4)
    heading('4.2.1 Phân trách nhiệm theo luồng đăng ký',5)
    code('@PostMapping("/register")\npublic ResponseEntity<UserResponseDTO> register(\n        @Valid @RequestBody CreateUserRequestDTO request) {\n    UserResponseDTO response = authService.register(request);\n    return ResponseEntity.status(HttpStatus.CREATED).body(response);\n}')
    para('AuthController nhận HTTP, ánh xạ JSON và kích hoạt Bean Validation; không thực hiện truy vấn hoặc hash mật khẩu. AuthService.register() trim dữ liệu định danh, kiểm tra trùng, hash mật khẩu và lưu tài khoản. TaiKhoanRepository thực hiện existsByUserName(), existsByEmail() và saveAndFlush(). UserResponseDTO được trả thay vì Entity, nên passwordHash không xuất hiện trong JSON.')
    para('Spring gọi phương thức Controller sau khi HandlerMapping/HandlerAdapter xác định handler và giải quyết tham số. @Valid lỗi sẽ chặn trước khi gọi service. Với API yêu cầu token, filter bảo mật chạy trước Controller; lỗi ở filter không tự đi qua @RestControllerAdvice.')
    heading('4.2.2 Validation và lỗi có cấu trúc',5)
    table(['Tình huống','HTTP','Xử lý'],[
        ('Thiếu/sai dữ liệu; JSON không đọc được','400','GlobalExceptionHandler; không trả rejected password'),
        ('Sai thông tin đăng nhập; token thiếu/sai/hết hạn','401','Handler MVC hoặc entry point Security, tùy nơi lỗi'),
        ('Username/email đã tồn tại','409','DuplicateAccountException; DB unique là lớp bảo vệ cuối'),
        ('Lỗi nội bộ không dự kiến','500','Response an toàn; chi tiết log phía server'),
    ],[5,1.5,9.4])
    para('Response có errorCode, errorMessage, details và timestamp theo ErrorResponseDTO. Timestamp hiện là LocalDateTime được tạo theo UTC; không được ghi rằng đã trả kiểu Instant có hậu tố Z. Mã 403 được cấu hình cho access denied nhưng không có nghĩa dự án đã hoàn thiện chính sách vai trò USER/ADMIN. Mã 404/quyền sở hữu dữ liệu chat thuộc luồng chat cần triển khai và kiểm thử riêng.')
    heading('4.2.3 Transaction và request cạnh tranh',5)
    para('AuthService.register() có @Transactional. Khi Controller gọi bean proxy, transaction bắt đầu trước nghiệp vụ và commit sau khi phương thức trả về thành công; runtime exception thoát ra sẽ rollback theo mặc định. saveAndFlush() đẩy SQL xuống DB để phát hiện lỗi sớm nhưng flush không đồng nghĩa commit. login()/me() dùng readOnly=true như gợi ý tối ưu, không coi đây là cơ chế phân quyền hoặc lệnh cấm ghi tuyệt đối.')
    para('Kiểm tra exists trước khi lưu không loại bỏ race condition. Nếu hai request cùng đăng ký, unique username/email trong PostgreSQL bảo vệ dữ liệu. Service bắt SQLSTATE 23505 tại saveAndFlush và chuyển thành DuplicateAccountException. Unit test concurrentUniqueConflictBecomesDuplicateAccount() mô phỏng lỗi này; chưa phải phép thử HTTP đồng thời trên database thật.')
    para('Cần bổ sung integration test gọi service qua proxy, tạo lỗi sau một bước ghi rồi kiểm tra DB không còn dữ liệu ghi dở; test phải chứng minh trạng thái database ngoài transaction thử nghiệm. Hiện không tuyên bố đã kiểm chứng rollback giữa nhiều lệnh ghi hoặc kết quả cạnh tranh đồng thời chỉ dựa trên mock.')
    heading('4.3 Spring Data JPA và lược đồ dữ liệu',4)
    heading('4.3.1 Entity và repository proxy',5)
    para('TaiKhoan dùng UUID, Instant và @CreationTimestamp; Entity khai báo @Table(name="tai_khoan") và các cột tương ứng migration. Conversation, Message, Notification dùng @ManyToOne(fetch=LAZY). Repository kế thừa JpaRepository<TaiKhoan, UUID>; findByUserName() được Spring Data phân tích thành truy vấn. Có thể quản lý quan hệ từ phía con mà không cần thêm collection @OneToMany vào mọi Entity.')
    code('public interface TaiKhoanRepository\n        extends JpaRepository<TaiKhoan, UUID> {\n    Optional<TaiKhoan> findByUserName(String userName);\n    boolean existsByUserName(String userName);\n    boolean existsByEmail(String email);\n}')
    heading('4.3.2 Migration và SQL thật',5)
    para('Flyway V1 tạo schema chat và unique constraint; V2 tạo dữ liệu mẫu; V3 bổ sung password_hash nếu chưa tồn tại. Code hiện V1 cũng có cột này nên V3 có thể ghi cảnh báo bỏ qua, không có nghĩa migration thất bại. ddl-auto=none: JPA không tự dựng schema thay migration. Không sửa nội dung migration đã chạy trên môi trường chia sẻ; thay đổi schema tiếp theo nên có version mới.')
    picture(sql_img,'Minh chứng MC03. Trích SQL Hibernate thật từ auth-startup.log, ngày 14/10/2026.')
    para('Hai truy vấn đầu kiểm tra username/email; INSERT ghi hash vào password_hash, không ghi mật khẩu rõ. Dấu ? là bind parameter; log này không hiển thị giá trị nhạy cảm. Truy vấn findByUserName dùng WHERE user_name=? khi đăng nhập. Không nối input người dùng vào chuỗi SQL. @CreationTimestamp do Hibernate điền lúc persist; seed chạy bằng SQL dùng default/timestamp trong migration.')
    heading('4.3.3 N+1 và cách kiểm chứng',5)
    para('Ví dụ rủi ro: lấy N Message rồi trong vòng lặp truy cập sender.tenTaiKhoan hoặc conversation có LAZY; nếu các quan hệ chưa được nạp, có thể phát sinh thêm query. N+1 không được chứng minh bằng việc nhìn annotation; cần ghi số SQL của cùng endpoint, cùng bộ dữ liệu trước/sau xử lý. Với luồng /me chỉ truy vấn một tài khoản, không suy ra rằng toàn bộ ứng dụng đã hết N+1.')
    para('Khi làm Chat REST, nhóm chọn fetch join, @EntityGraph hoặc DTO projection theo dữ liệu cần trả; không đổi tất cả quan hệ thành EAGER. Truy vấn lịch sử cần phân trang và thứ tự ổn định có ID phụ. Chưa có phép đo N+1 hoặc benchmark tối ưu trong phiên Authentication, nên bảng phủ kỹ thuật ghi trạng thái chưa đo.')
    heading('4.4 Auto configuration và starter trong Spring Boot 4',4)
    para('Phiên bản dùng là Spring Boot 4.1.1, Java 21. Thêm starter-data-jpa đưa các dependency JPA/JDBC/Hibernate cần thiết vào classpath, nhưng kết nối vẫn cần driver PostgreSQL và cấu hình datasource hợp lệ. Auto-configuration xem điều kiện về class, bean và property để quyết định tạo Hikari DataSource, EntityManagerFactory và transaction infrastructure. Cơ chế này không có nghĩa mọi bean luôn được tạo chỉ vì thêm starter.')
    para('Dự án khai báo riêng starter-webmvc, starter-data-jpa, starter-security, starter-oauth2-resource-server, starter-validation và starter-webmvc-test; Flyway có flyway-core cùng module PostgreSQL. Đó là dependency thực tế trong pom.xml. Chưa khai báo AMQP/SSE-specific infrastructure để coi RabbitMQ/SSE là chức năng đã hoàn thành.')
    code('.\\mvnw.cmd spring-boot:run "-Dspring-boot.run.arguments=--debug"\n# Tim Conditions Evaluation Report trong log.\n# Khong can them Actuator chi de xem cac dieu kien nay.')
    picture(cond_img,'Minh chứng MC05. Trích điều kiện matched trong auth-verify.log.')
    para('DataSourceAutoConfiguration matched vì có lớp JDBC và chưa có ConnectionFactory R2DBC; nhánh pooled tạo datasource khi chưa có DataSource tự khai báo. DataSourceConfiguration.Hikari matched vì có Hikari và không có bean DataSource khác. HibernateJpaAutoConfiguration matched khi có lớp JPA/Hibernate tương ứng. Đây là bằng chứng đã chạy; còn các ví dụ @ConditionalOnMissingBean khác phải đọc đúng điều kiện trong report, không suy từ tên annotation.')
    para('FlywayConfig hiện là cấu hình tự viết, có @ConditionalOnProperty và @Bean(initMethod="migrate"); cần phân biệt với auto-configuration Flyway của Spring Boot. Nhóm không ghi rằng chỉ nhờ starter là migration tự chạy mà bỏ qua cấu hình thực tế này.')
    heading('4.5 Kiểm thử và bằng chứng thực nghiệm',4)
    table(['Nhóm test','Cơ chế','Kết quả 14/10/2026'],[
        ('AuthServiceTest + AuthLoginTest','JUnit + Mockito, không khởi động Spring; BCrypt thật, repository/JWT mock','6 + 2 test thành công'),
        ('AuthControllerTest','MockMvc standalone, service mock; validation và exception advice','8 test thành công'),
        ('SecurityConfigTest','@WebMvcTest; chain/encoder/decoder thật, AuthService mock','11 test thành công'),
        ('NotificationDispatcherApplicationTests','@SpringBootTest, context và PostgreSQL thật','2 test thành công'),
    ],[4.5,7,4.4])
    picture(test_img,'Minh chứng MC04. Kết quả Maven clean verify từ log thực tế, không phải ảnh Postman.')
    para('Tổng cộng 29 test, Failures=0, Errors=0, BUILD SUCCESS. Unit test có ca đúng và lỗi; test slice xác minh JWT bằng bộ lọc thật nhưng không kiểm chứng SQL. Test SpringBootTest hiện kiểm tra context và repository count, chưa phải test rollback. Không gộp các loại này dưới tên “unit test toàn bộ hệ thống”.')

replace_section('Chương 4: Bắt buộc lõi của Spring Boot',
                'Chương 5 trở đi: Nội dung chọn thêm và thực nghiệm — linh hoạt theo bài toán',chapter4)
auth_anchor=anchor('Chương 5 trở đi: Nội dung chọn thêm và thực nghiệm — linh hoạt theo bài toán')
auth_anchor.text='Chương 5: Authentication và bảo vệ API'
auth_anchor.paragraph_format.page_break_before=True
anchor('Chương 4: Bắt buộc lõi của Spring Boot').paragraph_format.page_break_before=True

def authentication():
    para('Authentication là phần đã triển khai của backend chat 1–1. Cơ chế hiện tại gồm đăng ký, đăng nhập bằng mật khẩu và truy cập thông tin tài khoản qua JWT. SSE cùng RabbitMQ thuộc phần tích hợp tiếp theo; các nội dung này không được coi là đã kiểm thử chỉ vì Authentication hoạt động.')
    heading('5.1 Phạm vi và API contract',4)
    table(['Endpoint','Quyền','Kết quả và lỗi chính'],[
        ('POST /api/auth/register','Công khai','201 tạo tài khoản; 400 validation; 409 trùng username/email'),
        ('POST /api/auth/login','Công khai','200 token; 400 body/validation; 401 sai thông tin'),
        ('GET /api/users/me','Bearer JWT','200 tài khoản chủ token; 401 thiếu/sai/hết hạn token hoặc tài khoản không còn tồn tại'),
    ],[5,3.2,7.7])
    para('Request đăng ký dùng CreateUserRequestDTO với tenTaiKhoan, email, userName, password. LoginRequestDTO nhận userName/password. UserResponseDTO có id dạng chuỗi UUID, tenTaiKhoan, email, userName, createdAt. TokenResponseDTO có accessToken, tokenType và expiresIn (giây). Ví dụ contract dưới đây dùng placeholder, không phải token hoặc tài khoản thật:')
    code('POST /api/auth/login\nContent-Type: application/json\n{"userName":"<username>","password":"<password>"}\n\n200 OK\n{"accessToken":"<jwt>","tokenType":"Bearer","expiresIn":900}')
    heading('5.2 Đăng ký và bảo vệ mật khẩu',4)
    para('AuthService.register() dùng constructor injection và @Transactional. Username/email được trim; kiểm tra unique hiện phân biệt hoa/thường theo schema và query. Tên hiển thị được trim; mật khẩu giữ nguyên. PasswordEncoder bean là BCryptPasswordEncoder; encode() tạo hash và matches() kiểm chứng plaintext so với hash. Hai lần encode cùng mật khẩu có thể khác do salt, nên không so hai chuỗi hash vừa tạo.')
    para('BCrypt giới hạn 72 byte UTF-8. DTO hiện giới hạn theo số ký tự, service kiểm tra thêm byte UTF-8 để tránh cắt ngầm mật khẩu. Mật khẩu thô không được trả về DTO, token hoặc log nghiệp vụ. Ràng buộc UNIQUE PostgreSQL bảo vệ request cạnh tranh; chỉ SQLSTATE 23505 được ánh xạ sang trùng tài khoản, lỗi integrity khác không bị giả thành 409.')
    heading('5.3 Đăng nhập và cấp JWT',4)
    para('AuthService.login() tìm tài khoản bằng username, kiểm tra PasswordEncoder.matches() rồi gọi JwtService.issue(UUID). Username không tồn tại và sai mật khẩu đều trả cùng phản hồi 401. JWT được ký HS256 bằng khóa lấy từ JWT_SECRET: Base64 của ít nhất 32 byte ngẫu nhiên; không có khóa mặc định trong source. JWT_TTL_SECONDS mặc định 900 giây, giới hạn cấu hình 1 đến 86400 giây.')
    table(['Claim','Giá trị và vai trò'],[
        ('sub','UUID tài khoản; là nguồn danh tính cho /me'),
        ('iss','chat-system; decoder kiểm tra issuer'),
        ('iat, exp','Thời điểm cấp và hết hạn; decoder kiểm tra thời gian, không có clock skew trong cấu hình hiện tại'),
        ('jti','UUID nhận diện token; hiện chưa có kho thu hồi token'),
    ],[3,12.9])
    para('JwtService dùng JwtEncoder/Nimbus; JwtConfig cấu hình JwtDecoder với thuật toán HS256, timestamp validator, issuer validator và kiểm tra subject UUID cùng sự hiện diện iat/exp. Không dùng nội dung JWT chưa kiểm chứng để định danh user. Giai đoạn này chưa có refresh token, blacklist hoặc logout thu hồi token; token cũ có thể còn dùng đến khi hết hạn nếu khóa vẫn giữ nguyên.')
    heading('5.4 Xác thực request và ranh giới bảo mật',4)
    code('GET /api/users/me\nAuthorization: Bearer <accessToken>\n\nBearerTokenAuthenticationFilter -> JwtDecoder\n  -> Authentication / SecurityContext -> UserController\n  -> AuthService.me(UUID tu sub) -> TaiKhoanRepository')
    para('Spring Security OAuth2 Resource Server cung cấp bộ lọc Bearer; nhóm cấu hình và sử dụng cơ chế này, không tự viết một JwtFilter riêng. Filter kiểm chứng token trước khi vào Controller và đưa Authentication vào SecurityContext. UserController nhận @AuthenticationPrincipal Jwt; AuthService.me() truy vấn UUID từ sub. Query userId do client gửi không thay đổi danh tính trả về.')
    para('SecurityFilterChain chỉ cho POST /register và POST /login công khai, các request khác cần xác thực. Form login, HTTP Basic và session login bị tắt; policy STATELESS không lưu phiên đăng nhập giữa request. Entry point và access denied handler trả cùng cấu trúc JSON lỗi thay vì redirect HTML. AuthenticationException trong MVC được GlobalExceptionHandler xử lý riêng.')
    para('Cấu hình CSRF tắt vì xác thực hiện chỉ nhận Bearer trong header, không dùng cookie tự đính kèm. Không suy ra stateless tự động làm mọi ứng dụng an toàn trước CSRF: nếu chuyển sang cookie/session cần đánh giá và cấu hình lại. Chưa cấu hình CORS mở rộng; frontend khác origin phải có allowlist cụ thể. Hiện các API yêu cầu authenticated, chưa triển khai phân vai ADMIN/USER hoặc quyền truy cập hội thoại.')
    para('SSE dự kiến phải dùng danh tính đã xác thực để chọn kênh người nhận. Nhóm sẽ chọn cơ chế client truyền token, kiểm tra token hết hạn, hủy kết nối và chống đăng ký kênh user khác; chưa có endpoint SSE trong phiên bản Authentication này. Không truyền access token vào URL mà không đánh giá rủi ro log và cơ chế bảo vệ.')
    heading('5.5 Kết quả kiểm thử API và Security',4)
    para('Bảng dưới là biên bản phiên chạy collection Postman bằng Newman 6.2.1 ngày 14/10/2026. Collection tạo dữ liệu ngẫu nhiên trên PostgreSQL test và dùng token của lần đăng nhập; không dùng password seed được suy đoán từ comment. Database tạm đã bị xóa sau kiểm thử, nên các UUID/token của phiên chạy không được lưu trong báo cáo.')
    table(['Ca kiểm thử','HTTP quan sát','Ý nghĩa'],[
        ('Đăng ký hợp lệ','201','Tạo tài khoản; response không có password/passwordHash'),
        ('Đăng ký lại tài khoản','409','Chặn trùng username/email'),
        ('Input không hợp lệ','400','Validation chặn request'),
        ('Đăng nhập đúng','200','Nhận Bearer token và expiresIn > 0'),
        ('/me có token hợp lệ','200','ID khớp user đã đăng ký'),
        ('/me thêm query userId khác','200','Vẫn trả chủ token, không đổi danh tính'),
        ('/me thiếu token','401','Không truy cập ẩn danh'),
        ('/me token lỗi','401','Token không hợp lệ bị từ chối'),
        ('Đăng nhập sai mật khẩu','401','Không cấp token'),
    ],[6,2.5,7.4])
    para('Kết quả: 9 request và 11 assertion đều thành công. Đây là kiểm thử HTTP có database thật qua Newman (runner cho Postman collection), không phải ảnh chụp giao diện Postman. Collection: docs/postman/authentication.postman_collection.json. Các ca JWT hết hạn, sai chữ ký và sai issuer còn được kiểm tra bằng SecurityConfigTest với encoder/decoder thật. Chưa có số liệu latency lặp đủ để công bố p95 hoặc năng lực chịu tải.')
    code('npx.cmd --yes newman@6.2.1 run\n  docs/postman/authentication.postman_collection.json\n  --env-var baseUrl=http://localhost:18081\n# Lenh tren trinh bay nhieu dong de de doc; khi chay ghep cung dong.')
    heading('5.6 Cấu hình bí mật và phần việc còn lại',4)
    para('Profile dev đọc .env local; file mẫu chỉ có placeholder cho secret. Supabase URL, username và password được lấy qua biến môi trường, đã bỏ credential dự phòng trong source. Việc xóa khỏi source không thu hồi credential đã lộ trước đó: người quản lý Supabase phải xác nhận đã đổi mật khẩu và cập nhật môi trường/CI; chưa có bằng chứng thu hồi trong bộ kết quả này.')
    para('Các giới hạn còn lại gồm integration test rollback và request đăng ký đồng thời thật, kiểm thử quyền sở hữu hội thoại, chính sách giới hạn đăng nhập và tích hợp SSE. Tìm kiếm bạn bè, chat nhóm, refresh token và đăng nhập mạng xã hội không nằm trong phần Authentication đã nghiệm thu ở đây.')

insert_block_before(anchor('Chương 6'),authentication)

def coverage():
    para(f'Bảng phủ kỹ thuật đối chiếu code commit {commit} và kết quả thực nghiệm ngày 14/10/2026. Đường dẫn Java tính từ notification-dispatcher/src/main/java/com/notification/dispatcher/; đường dẫn test từ src/test/java/com/notification/dispatcher/. “Có code” không đồng nghĩa đã kiểm chứng tất cả hành vi.')
    table(['Kỹ thuật','File hoặc cấu hình','Minh chứng','Trạng thái'],[
        ('IoC, constructor DI','service/AuthService.java; config/SecurityConfig.java','Đoạn constructor 4.1; test context','Có code và context chạy'),
        ('Scope, @Primary/@Qualifier','Giải thích 4.1.1–4.1.2','Ví dụ chọn bean, chưa bật hai encoder','Giải thích; chưa demo hai bean'),
        ('Vòng đời bean','config/FlywayConfig.java','initMethod=migrate; MC02','Đã thấy init; chưa log PostConstruct/Destroy'),
        ('Auto-configuration','pom.xml; app datasource','MC05, Conditions Evaluation Report','Đã có điều kiện matched'),
        ('REST và DTO','controller/AuthController.java; DTO/*','POST register; AuthControllerTest','Đã kiểm thử'),
        ('JPA, query method','entity/TaiKhoan.java; repository/TaiKhoanRepository.java','MC03; query username/email và INSERT','Đã quan sát SQL'),
        ('Migration','resources/db/migration/V1–V3','MC02; schema mới lên v3','Đã kiểm thử DB mới'),
        ('Transaction','service/AuthService.java','@Transactional; mock lỗi lưu','Có code; chưa chứng minh rollback nhiều lệnh'),
        ('Unique và cạnh tranh','V1; AuthService.saveAndFlush','409 trong Newman; mock SQLSTATE 23505','Chưa chạy hai request đồng thời thật'),
        ('BCrypt','SecurityConfig; AuthService','AuthServiceTest/AuthLoginTest','Đã kiểm thử đúng/sai'),
        ('JWT và SecurityContext','JwtConfig; security/JwtService; UserController','SecurityConfigTest; /me Newman','Đã kiểm thử chữ ký/issuer/expiry/identity'),
        ('Lỗi thống nhất','exception/GlobalExceptionHandler; SecurityConfig','400/401/409 và test 500 an toàn','Đã kiểm thử các ca đã liệt kê'),
        ('JUnit, Mockito, MVC slice','service/*Test; controller/*Test; config/*Test','MC04: 29 test toàn bộ','Đã chạy; phân loại ở 4.5'),
        ('Profile và secret','application*.properties; .env.example','Không có credential dự phòng Supabase','Có cấu hình; thu hồi secret cũ chờ xác nhận'),
        ('N+1','LAZY associations; thiết kế 4.3.3','Chưa có đo trước/sau trên Chat API','Chưa kiểm chứng'),
        ('SSE và RabbitMQ','Kế hoạch chat 1–1','Chưa có endpoint/worker trong code đối chiếu','Chưa triển khai'),
        ('CI, Compose toàn hệ thống','Chưa thấy workflow CI/app Docker trên nhánh đối chiếu; compose hiện có DB','Không dùng kết quả nhánh cũ thay bằng chứng hiện tại','Cần tích hợp và kiểm chứng lại'),
    ],[2.8,4.8,4.9,3.4])
    para('Minh chứng tái hiện: MC01 công cụ; MC02 database/migration/khởi động; MC03 SQL; MC04 Maven test; MC05 auto-configuration; biên bản Newman ở 5.5. Không đánh dấu các ô chưa kiểm chứng là hoàn thành khi chưa có test hoặc log tương ứng. Log có thể bị .gitignore loại, nên khi nộp cần lưu bản đã che thông tin nhạy cảm cùng phụ lục hoặc artifact của phiên chạy.')

replace_section('12.3. Bảng phủ kỹ thuật','PHẦN 4: TỔNG KẾT VÀ BÀI HỌC KINH NGHIỆM',coverage)

# Ensure new headings stay black, preserve existing layout and images elsewhere.
anchor('12.3. Bảng phủ kỹ thuật').paragraph_format.page_break_before=True
# Keep the existing cover design but prevent a separate page containing only its date.
cover_date=next(p for p in doc.paragraphs if p.text.startswith('Thành phố Hồ Chí Minh, tháng'))
previous=cover_date._p.getprevious()
for unused in range(3):
    if previous is None or previous.tag!=qn('w:p') or ''.join(previous.itertext()).strip():break
    older=previous.getprevious();previous.getparent().remove(previous);previous=older
for cell in doc.tables[0].rows[0].cells:
    pass
for row in doc.tables[0].rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:r.font.size=Pt(12)
for p in doc.paragraphs:
    if p.text.startswith('Hình 1.3. Khởi tạo dự án Spring Boot'):
        for r in p.runs:r.text=r.text.replace('Hình 1.3.','Hình 3.1.')
for st in ['Heading 3','Heading 4','Heading 5']:
    if st in doc.styles:doc.styles[st].font.color.rgb=RGBColor(0,0,0)
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
out=OUT/'Bao_cao_do_an_bo_sung_Authentication.docx'
doc.save(out)
print(out)
print(json.dumps({'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'inline_images':len(doc.inline_shapes),'commit':commit}))
